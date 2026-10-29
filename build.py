#!/usr/bin/env python3
"""Сборщик каталога номеров СоцГума.

Берёт данные из Google Таблицы (или из демо-файла), проверяет их, готовит обложки
и собирает готовый статический сайт в папку site/.

Если в таблице есть ошибки, сайт НЕ пересобирается: на публикации остаётся прошлая
версия, а в отчёте написано, на каком листе и в какой строке что исправить.

Переменные окружения (все необязательные):
  SHEET_ID        идентификатор Google Таблицы (длинный код из её адреса)
  SHEET_XLSX_URL  прямая ссылка на xlsx (вместо SHEET_ID)
  SHEET_FILE      путь к локальному xlsx-файлу (для проверки на своём компьютере)
Если не задано ничего, используется data/demo.xlsx и на сайте показывается пометка «демо».
"""
import datetime as dt
import difflib
import hashlib
import html
import io
import json
import os
import random
import re
import shutil
import sys
import tempfile
import time
import urllib.error
import urllib.request
import zipfile
from collections import OrderedDict, defaultdict
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import openpyxl
from jinja2 import Environment, FileSystemLoader, select_autoescape
from PIL import Image, ImageDraw, ImageFilter, ImageOps

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "site"
TMP = ROOT / "site.tmp"
COVERS = ROOT / "covers"
Image.MAX_IMAGE_PIXELS = 200_000_000

# ---------------------------------------------------------------- вспомогательное

TRANSLIT = dict(zip("абвгдежзийклмнопрстуфхцчшщъыьэюя", [
    "a", "b", "v", "g", "d", "e", "zh", "z", "i", "y", "k", "l", "m", "n", "o", "p", "r", "s", "t",
    "u", "f", "kh", "ts", "ch", "sh", "shch", "", "y", "", "e", "yu", "ya"]))


def norm(s):
    """Для сравнения имён и названий: без регистра, без лишних пробелов, ё = е."""
    return re.sub(r"\s+", " ", str(s or "").strip()).casefold().replace("ё", "е")


def slugify(s):
    s = norm(s)
    s = "".join(TRANSLIT.get(ch, ch) for ch in s)
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s or "x"


class Slugs:
    def __init__(self):
        self.used = set()

    def make(self, text):
        base = slugify(text)
        slug, i = base, 2
        while slug in self.used:
            slug = f"{base}-{i}"
            i += 1
        self.used.add(slug)
        return slug


ID_RE = re.compile(r"[A-Za-z0-9_-]+")


def assign_slugs(rows, sheet, report, label):
    """Адрес строки = её ID из таблицы. Если ID нет, делаем из названия и предупреждаем:
    такой адрес сменится при переименовании. Возвращает ({номер строки: адрес}, Slugs)."""
    slugs, out = Slugs(), {}
    for r in rows:
        sid = r.get("id", "")
        if not sid:
            continue
        if not ID_RE.fullmatch(sid):
            report.error(sheet, r["_row"], f"ID «{sid}» должен состоять только из латинских букв, цифр, дефиса и подчёркивания.")
        elif sid.casefold() in slugs.used:
            report.error(sheet, r["_row"], f"ID «{sid}» уже встречался выше. У каждой строки свой ID.")
        else:
            slugs.used.add(sid.casefold())
            out[r["_row"]] = sid
    missing = []
    for r in rows:
        if r["_row"] not in out and r.get("name"):
            out[r["_row"]] = slugs.make(r["name"])
            missing.append(r["_row"])
    if missing:
        report.warn(sheet, missing[0], f"{label}: у {len(missing)} строк нет ID, адрес взят из названия и сменится при переименовании. "
                                       "Выберите в таблице Каталог → Проставить ID.")
    return out, slugs


def hms(sec):
    sec = int(sec)
    h, rem = divmod(sec, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def vk_t(sec):
    sec = int(sec)
    h, rem = divmod(sec, 3600)
    m, s = divmod(rem, 60)
    out = (f"{h}h" if h else "") + (f"{m}m" if m or h else "") + f"{s}s"
    return out


MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
          "сентября", "октября", "ноября", "декабря"]


def ru_date(d):
    return f"{d.day} {MONTHS[d.month - 1]} {d.year}" if d else ""


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, dt.time):
        return f"{v.hour}:{v.minute:02d}:{v.second:02d}"
    if isinstance(v, dt.timedelta):
        return hms(v.total_seconds())
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_seconds(v):
    """Время из таблицы -> секунды. None, если не удалось разобрать."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.time):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, dt.timedelta):
        return int(v.total_seconds())
    if isinstance(v, dt.datetime):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, (int, float)):
        if 0 < v < 1:  # доля суток, если ячейка в формате «время»
            return int(round(v * 86400))
        return int(v)
    s = str(v).strip().lower().replace("ч", "h").replace("м", "m").replace("с", "s")
    m = re.fullmatch(r"(?:(\d+):)?(\d{1,2}):(\d{2})(?:[.,]\d+)?", s)
    if m:
        h, mi, se = int(m.group(1) or 0), int(m.group(2)), int(m.group(3))
        if mi > 59 or se > 59:
            return None
        return h * 3600 + mi * 60 + se
    m = re.fullmatch(r"(?:(\d+)h)?\s*(?:(\d+)m)?\s*(?:(\d+)s)?", s)
    if m and any(m.groups()):
        return int(m.group(1) or 0) * 3600 + int(m.group(2) or 0) * 60 + int(m.group(3) or 0)
    if re.fullmatch(r"\d+", s):
        return int(s)
    return None


def parse_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%Y.%m.%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def parse_video(url):
    """Ссылка VK -> параметры встраивания. None, если это не ссылка на видео VK."""
    url = html.unescape(str(url or "")).strip()
    if not url:
        return None
    parsed = urlparse(url if "://" in url else "https://" + url)
    q = parse_qs(parsed.query)
    oid = vid = None
    h = (q.get("hash") or [""])[0]
    if "oid" in q and "id" in q:
        oid, vid = q["oid"][0], q["id"][0]
    else:
        m = re.search(r"video(-?\d+)_(\d+)", url)
        if m:
            oid, vid = m.group(1), m.group(2)
    if not (oid and vid and re.fullmatch(r"-?\d+", oid) and re.fullmatch(r"\d+", vid)):
        return None
    host = "vkvideo.ru" if "vkvideo.ru" in parsed.netloc else "vk.com"
    if h and not re.fullmatch(r"[0-9a-fA-F]+", h):
        h = ""
    return {"oid": oid, "id": vid, "hash": h, "host": host}


def season_of(raw, date, year):
    """Учебный год: «25-26» или «2025–2026» из таблицы, иначе по дате (с сентября начинается новый год).
    Возвращает (подпись «2025–2026», год начала) или ("", 0)."""
    m = re.fullmatch(r"\s*(\d{2,4})\s*[-–—/]\s*(\d{2,4})\s*", str(raw or ""))
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        a = a + 2000 if a < 100 else a
        b = b + 2000 if b < 100 else b
        if b == a + 1:
            return f"{a}–{b}", a
    if date:
        a = date.year if date.month >= 9 else date.year - 1
        return f"{a}–{a + 1}", a
    return "", 0


def watch_link(url):
    """Ссылка на запись, которую нельзя встроить плеером (пост VK, YouTube): кнопка «Смотреть…»."""
    url = html.unescape(str(url or "")).strip()
    if not re.match(r"^https?://\S+$", url, re.I):
        return None
    host = urlparse(url).netloc.lower()
    if host.endswith(("vk.com", "vk.ru", "vkvideo.ru")):
        label = "во ВКонтакте"
    elif "youtube.com" in host or "youtu.be" in host:
        label = "на YouTube"
    else:
        label = "по ссылке"
    return {"url": url, "label": label}


SOCIAL_KINDS = [
    (("vk.com", "vk.ru", "vkvideo.ru"), "vk", "ВКонтакте"),
    (("t.me", "telegram.me", "telegram.org"), "tg", "Telegram"),
    (("instagram.com",), "ig", "Instagram"),
    (("youtube.com", "youtu.be"), "yt", "YouTube"),
    (("tiktok.com",), "tt", "TikTok"),
    (("rutube.ru",), "rt", "Rutube"),
]


def social_links(text):
    """Ссылки на группы и соцсети из ячейки (через пробел, запятую или перенос строки): значок и название сети."""
    out, seen = [], set()
    for part in re.split(r"[\s,;]+", html.unescape(str(text or ""))):
        url = part.strip()
        if not re.match(r"^https?://\S+$", url, re.I) or url in seen:
            continue
        seen.add(url)
        host = urlparse(url).netloc.lower().removeprefix("www.")
        kind, label = "web", host or "Сайт"
        for hosts, k, name in SOCIAL_KINDS:
            if any(host == h or host.endswith("." + h) for h in hosts):
                kind, label = k, name
                break
        out.append({"url": url, "kind": kind, "label": label})
    return out


def split_people(text):
    seen, out = set(), []
    for part in re.split(r"[,;\n]+", text or ""):
        name = re.sub(r"\s+", " ", part).strip()
        if name and norm(name) not in seen:
            seen.add(norm(name))
            out.append(name)
    return out


# ---------------------------------------------------------------- отчёт

class Report:
    def __init__(self):
        self.errors, self.warnings = [], []

    @staticmethod
    def _where(sheet, row):
        return f"лист «{sheet}», строка {row}" if row else f"лист «{sheet}»"

    def error(self, sheet, row, text):
        self.errors.append((self._where(sheet, row), text))

    def warn(self, sheet, row, text):
        self.warnings.append((self._where(sheet, row), text))

    def markdown(self, source, stats):
        lines = ["# Отчёт сборки каталога", "", f"Источник данных: {source}", ""]
        if stats:
            lines += [stats, ""]
        if self.errors:
            lines += [f"## Ошибки: {len(self.errors)}. Сайт не обновлён, на публикации осталась прошлая версия.", ""]
            lines += [f"- {w}: {t}" for w, t in self.errors] + [""]
        else:
            lines += ["## Ошибок нет", ""]
        if self.warnings:
            lines += [f"## Предупреждения: {len(self.warnings)}. Сайт собран, но стоит проверить.", ""]
            lines += [f"- {w}: {t}" for w, t in self.warnings] + [""]
        return "\n".join(lines)


# ---------------------------------------------------------------- чтение таблицы

SPEC = {
    "Концерты": {
        "name": ["название"], "date": ["дата"], "video": ["видео", "ссылка на видео"], "about": ["описание"],
        "id": ["id"], "season": ["учебный год", "поколение"], "year": ["год"], "topic": ["тема"],
        "cover": ["обложка"], "video2": ["видео 2", "видео2"],
    },
    "Направления": {"name": ["название"], "about": ["описание"], "id": ["id"], "kind": ["тип", "вне сцены"],
                    "c1": ["критерий 1"], "c2": ["критерий 2"], "c3": ["критерий 3"], "c4": ["критерий 4"],
                    "c5": ["общее впечатление"], "rnote": ["пояснение к оценке", "пояснение"]},
    "Номера": {
        "id": ["id"], "concert": ["концерт"], "direction": ["направление"],
        "order": ["№ в концерте", "номер в концерте", "порядок"], "title": ["название"],
        "collective": ["коллектив"], "performers": ["участники", "исполнители"], "choreographers": ["постановщики"],
        "about": ["описание"], "cover": ["обложка"], "start": ["начало"], "end": ["конец"],
        "old_url": ["старый адрес"], "links": ["ссылки", "соцсети", "группа"], "direction2": ["ещё направление", "еще направление", "второе направление"],
    },
    "Люди": {"name": ["имя"], "about": ["о себе", "описание"], "link": ["ссылка"], "id": ["id"]},
    "Формы": {"title": ["форма"], "url": ["ссылка для сотрудников", "ссылка"]},
}
REQUIRED = {
    "Концерты": ["name"], "Направления": ["name"],
    "Номера": ["id", "concert", "title"], "Люди": ["name"], "Формы": ["title", "url"],
}
OPTIONAL_SHEETS = {"Люди", "Формы"}
RAW_COLS = {("Номера", "start"), ("Номера", "end")}  # время разбираем сами


FORMS = []   # ссылки на формы из листа «Формы», попадают в подвал сайта


def hnorm(s):
    return re.sub(r"[^0-9a-zа-я№]+", "", norm(s))


def open_workbook(report):
    """Возвращает (workbook, описание источника, демо?)."""
    local = os.environ.get("SHEET_FILE", "").strip()
    url = os.environ.get("SHEET_XLSX_URL", "").strip()
    sheet_id = os.environ.get("SHEET_ID", "").strip()
    if sheet_id and not url:
        if not re.fullmatch(r"[A-Za-z0-9_-]{20,}", sheet_id):
            report.error("настройки", None, "SHEET_ID похож не на идентификатор таблицы. Нужен длинный код из адреса таблицы между /d/ и /edit.")
            return None, "", False
        url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    if local:
        return openpyxl.load_workbook(local, data_only=True), f"файл {local}", False
    if url:
        data = None
        for attempt in range(3):
            try:
                req = urllib.request.Request(url, headers={"User-Agent": "socgum-catalog-builder/1.0"})
                with urllib.request.urlopen(req, timeout=60) as r:
                    data = r.read()
                break
            except (urllib.error.URLError, TimeoutError) as e:
                err = str(e)
                time.sleep(3 * (attempt + 1))
        if data is None:
            report.error("настройки", None, f"Не удалось скачать таблицу ({err}). Проверьте доступ: «Все, у кого есть ссылка — Читатель».")
            return None, "", False
        if not zipfile.is_zipfile(io.BytesIO(data)):
            report.error("настройки", None, "Вместо таблицы пришла веб-страница. Откройте доступ «Все, у кого есть ссылка — Читатель» и проверьте SHEET_ID.")
            return None, "", False
        return openpyxl.load_workbook(io.BytesIO(data), data_only=True), "Google Таблица", False
    return openpyxl.load_workbook(ROOT / "data" / "demo.xlsx", data_only=True), "демо-данные (data/demo.xlsx)", True


def read_sheet(wb, name, report):
    """Список словарей {поле: значение, '_row': номер строки}."""
    titles = {norm(t): t for t in wb.sheetnames}
    if norm(name) not in titles:
        if name not in OPTIONAL_SHEETS:
            report.error(name, None, f"Нет листа «{name}». Не переименовывайте листы.")
        return []
    ws = wb[titles[norm(name)]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        report.error(name, 1, "Лист пуст: нет строки с заголовками.")
        return []
    pos = {}
    hdr = {hnorm(h): i for i, h in enumerate(header) if h is not None}
    for field, aliases in SPEC[name].items():
        for a in aliases:
            if hnorm(a) in hdr:
                pos[field] = hdr[hnorm(a)]
                break
    for field in REQUIRED[name]:
        if field not in pos:
            report.error(name, 1, f"Нет столбца «{SPEC[name][field][0]}». Не меняйте заголовки в первой строке.")
    if any(f not in pos for f in REQUIRED[name]):
        return []
    out = []
    for i, row in enumerate(rows, start=2):
        rec = {"_row": i}
        for field, idx in pos.items():
            v = row[idx] if idx < len(row) else None
            rec[field] = v if (name, field) in RAW_COLS or field == "date" else cell_text(v)
        if all(cell_text(rec.get(f)) == "" for f in pos):
            continue
        out.append(rec)
    return out


# ---------------------------------------------------------------- модель данных

FORM_URL_RE = re.compile(r"^https://(docs\.google\.com/forms/|forms\.gle/)[^\s]+$")


def form_links(rows):
    """Лист «Формы» -> три кнопки для подвала: добавить концерт, добавить номер, дополнить номер."""
    found = {}
    for r in rows:
        title, url = norm(r.get("title") or ""), (r.get("url") or "").strip()
        if not FORM_URL_RE.match(url):
            continue
        if "дополн" in title:
            found.setdefault("C", url)
        elif "концерт" in title and "номер" not in title:
            found.setdefault("A", url)
        elif "номер" in title:
            found.setdefault("B", url)
    labels = {"A": "Добавить концерт", "B": "Добавить номер", "C": "Дополнить номер"}
    return [{"label": labels[k], "url": found[k]} for k in ("A", "B", "C") if k in found]


# Критерии оценки. Ключи в базе всегда одни и те же (idea, comp, perf, stage, overall), меняются только названия и подсказки.
# Формат: (ключ, название, подсказка, короткое название для строки под комментарием)
RATE_KEYS = ["idea", "comp", "perf", "stage", "overall"]
RATE_STAGE = [
    ("idea", "Замысел", "идея, образы, история номера", "замысел"),
    ("comp", "Постановка", "что придумал постановщик: рисунок, переходы, ритм, музыка", "постановка"),
    ("perf", "Исполнение", "как выступили артисты: техника, синхронность, игра", "исполнение"),
    ("stage", "Сценическое воплощение", "свет, костюмы, видео, реквизит", "сцена"),
    ("overall", "Общее впечатление", "эмоции, атмосфера, вайб", "впечатление"),
]
# Запасные критерии для направлений вне сцены. Их можно переопределить в листе «Направления» (столбцы «Критерий 1» … «Общее впечатление»).
RATE_OFF = {
    "админка": [
        ("Организованность", "порядок, сроки, чёткие договорённости", "организованность"),
        ("Понятность и связь", "ясно ли, что, где и когда; быстро ли отвечали", "связь"),
        ("Забота об участниках", "внимание, поддержка, атмосфера в команде", "забота"),
        ("Итог для концерта", "зал, расписание, всё ли сложилось в день концерта", "итог"),
        ("Общее впечатление", "чувствовали ли вы, что о вас позаботились", "впечатление"),
    ],
    "smm": [
        ("Подача и стиль", "оформление, единый визуальный язык", "подача"),
        ("Идея и тексты", "рубрики, анонсы, как рассказали о концерте", "тексты"),
        ("Живость и отклик", "интерес людей, ощущение живой группы", "отклик"),
        ("Фото и видео", "качество съёмки, монтажа, фотоотчётов", "фото и видео"),
        ("Общее впечатление", "захотелось ли прийти и участвовать", "впечатление"),
    ],
    "сценарка": [
        ("История и идея", "о чём концерт и как это придумано", "идея"),
        ("Тексты и диалоги", "слова ведущих и переходов", "тексты"),
        ("Связность концерта", "как номера собрались в одну историю", "связность"),
        ("Юмор и эмоции", "где улыбнулись и где задело", "юмор"),
        ("Общее впечатление", "как сценарий работал на концерт", "впечатление"),
    ],
    "декор": [
        ("Замысел", "образ и идея оформления", "замысел"),
        ("Исполнение на сцене", "как выглядит вживую: качество, детали", "исполнение"),
        ("Вписанность в номера", "помогал ли декор номерам и сцене", "вписанность"),
        ("Аккуратность", "надёжность, чистота работы", "аккуратность"),
        ("Общее впечатление", "создал ли декор атмосферу", "впечатление"),
    ],
    "рабы сцены": [
        ("Незаметность", "смены проходили без лишнего шума и пауз", "незаметность"),
        ("Чёткость и скорость", "реквизит, свет и смены вовремя", "чёткость"),
        ("Слаженность команды", "каждый знал своё место", "слаженность"),
        ("Безопасность", "аккуратность с реквизитом и техникой", "безопасность"),
        ("Общее впечатление", "всё ли шло как по маслу", "впечатление"),
    ],
}
RATE_OFF_GENERIC = [
    ("Организованность", "порядок и чёткость работы", "организованность"),
    ("Результат", "что получилось в итоге", "результат"),
    ("Команда", "слаженность и атмосфера", "команда"),
    ("Польза для концерта", "насколько это помогло концерту", "польза"),
    ("Общее впечатление", "ощущение от работы направления в целом", "впечатление"),
]


def _split_crit(cell):
    """«Название — подсказка» → (название, подсказка)."""
    t = (cell or "").strip()
    for sep in (" — ", " – ", " - "):
        if sep in t:
            a, b = t.split(sep, 1)
            return a.strip(), b.strip()
    return t, ""


def rate_criteria(r, offstage):
    """Пять критериев направления: из листа «Направления», иначе запасные. None — для сцены по умолчанию."""
    cells = [r.get(k) for k in ("c1", "c2", "c3", "c4", "c5")]
    given = [c for c in cells if (c or "").strip()]
    if not given and not offstage:
        return None
    base = RATE_OFF.get(norm(r.get("name", "")), RATE_OFF_GENERIC) if offstage else [(k[1], k[2], k[3]) for k in RATE_STAGE]
    out = []
    for i, key in enumerate(RATE_KEYS):
        if (cells[i] or "").strip():
            name, hint = _split_crit(cells[i])
            out.append((key, name, hint, name.lower()))
        else:
            b = base[i]
            out.append((key, b[0], b[1], b[2]))
    return out



def build_model(wb, report):
    sheets = {n: read_sheet(wb, n, report) for n in SPEC}
    FORMS[:] = form_links(sheets.get("Формы") or [])
    c_slug, _ = assign_slugs(sheets["Концерты"], "Концерты", report, "Концерты")
    d_slug, _ = assign_slugs(sheets["Направления"], "Направления", report, "Направления")
    p_slug, slugs_p = assign_slugs(sheets["Люди"], "Люди", report, "Люди")

    # Концерты
    concerts = OrderedDict()
    for r in sheets["Концерты"]:
        name = r["name"]
        if not name:
            report.error("Концерты", r["_row"], "Не заполнено название.")
            continue
        if norm(name) in concerts:
            report.error("Концерты", r["_row"], f"Концерт «{name}» записан дважды.")
            continue
        date = parse_date(r.get("date"))
        if r.get("date") not in (None, "") and not date:
            report.warn("Концерты", r["_row"], f"Дата «{cell_text(r['date'])}» не распознана. Формат: ГГГГ-ММ-ДД.")
        video_raw = r.get("video", "")
        video = parse_video(video_raw)
        links = []
        for raw in (video_raw if not video else "", r.get("video2", "")):
            if raw:
                wl = watch_link(raw)
                if wl:
                    links.append(wl)
                else:
                    report.warn("Концерты", r["_row"], f"Ссылка «{raw}» не похожа на ссылку на запись. Кнопки на сайте не будет.")
        season, season_start = season_of(r.get("season"), date, r.get("year"))
        if r.get("season") and not season:
            report.warn("Концерты", r["_row"], f"Учебный год «{r['season']}» не распознан. Пишите так: 25-26 или 2025–2026.")
        year = int(r["year"]) if re.fullmatch(r"\d{4}", r.get("year") or "") else (date.year if date else 0)
        cover_url = r.get("cover", "") if re.match(r"^https?://\S+$", r.get("cover", "") or "") else ""
        concerts[norm(name)] = {
            "name": name, "slug": c_slug[r["_row"]], "date": date, "video": video, "links": links, "row": r["_row"],
            "about": r.get("about", ""), "numbers": [], "season": season, "season_start": season_start,
            "year": year, "topic": r.get("topic", ""), "cover_url": cover_url,
        }

    # Направления
    directions = OrderedDict()
    for r in sheets["Направления"]:
        name = r["name"]
        if name and norm(name) not in directions:
            directions[norm(name)] = {"name": name, "slug": d_slug[r["_row"]], "about": r.get("about", ""),
                                      "offstage": "вне" in norm(r.get("kind", "")),
                                      "criteria": rate_criteria(r, "вне" in norm(r.get("kind", ""))),
                                      "rate_note": (r.get("rnote") or "").strip() or None}

    def suggest(value, options):
        close = difflib.get_close_matches(norm(value), list(options), n=1, cutoff=0.6)
        return f" Возможно, имелось в виду «{options[close[0]]['name']}»." if close else ""

    # Номера
    numbers, ids = OrderedDict(), set()
    for r in sheets["Номера"]:
        row, nid = r["_row"], r["id"]
        ok = True
        if not nid:
            report.error("Номера", row, "Не заполнен ID."); ok = False
        elif not re.fullmatch(r"[A-Za-z0-9_-]+", nid):
            report.error("Номера", row, f"ID «{nid}» должен состоять только из латинских букв, цифр, дефиса и подчёркивания."); ok = False
        elif nid.casefold() in ids:
            report.error("Номера", row, f"ID «{nid}» уже встречался выше. Каждый номер должен иметь свой ID."); ok = False
        if not r["title"]:
            report.error("Номера", row, "Не заполнено название."); ok = False
        concert = concerts.get(norm(r["concert"]))
        if not r["concert"]:
            report.error("Номера", row, "Не выбран концерт."); ok = False
        elif not concert:
            report.error("Номера", row, f"Концерта «{r['concert']}» нет на листе «Концерты».{suggest(r['concert'], concerts)}"); ok = False
        dirs_found = []
        for cell in (r.get("direction"), r.get("direction2")):
            if not cell:
                continue
            whole = directions.get(norm(cell))
            parts = [cell] if whole else [x.strip() for x in re.split(r"[;,/]+", cell) if x.strip()]
            for part in parts:
                d = directions.get(norm(part))
                if not d:
                    report.error("Номера", row, f"Направления «{part}» нет на листе «Направления».{suggest(part, directions)}"); ok = False
                elif d not in dirs_found:
                    dirs_found.append(d)
        direction = dirs_found[0] if dirs_found else None
        if not dirs_found and ok:
            report.warn("Номера", row, f"У номера «{r['title']}» не выбрано направление: он не попадёт в фильтр по направлениям.")
        start, end = parse_seconds(r.get("start")), parse_seconds(r.get("end"))
        if cell_text(r.get("start")) and start is None:
            report.error("Номера", row, f"Начало «{cell_text(r['start'])}» не распознано. Формат: ч:мм:сс, например 0:20:52."); ok = False
        if cell_text(r.get("end")) and end is None:
            report.error("Номера", row, f"Конец «{cell_text(r['end'])}» не распознан. Формат: ч:мм:сс, например 0:24:10."); ok = False
        if start is not None and end is not None:
            if end <= start:
                report.error("Номера", row, f"Конец ({hms(end)}) раньше начала ({hms(start)})."); ok = False
            elif end - start > 30 * 60:
                report.warn("Номера", row, f"Номер «{r['title']}» длится {hms(end - start)}. Если это ошибка, проверьте формат времени (ч:мм:сс).")
        order = None
        if r.get("order"):
            if re.fullmatch(r"\d+", r["order"]):
                order = int(r["order"])
            else:
                report.warn("Номера", row, f"№ в концерте «{r['order']}» не число, порядок возьмётся по времени начала.")
        if not ok:
            continue
        ids.add(nid.casefold())
        num = {
            "id": nid, "row": row, "title": r["title"], "collective": r.get("collective", ""),
            "concert": concert, "direction": direction, "directions": dirs_found, "order": order,
            "performers": split_people(r.get("performers")), "choreographers": split_people(r.get("choreographers")),
            "about": r.get("about", ""), "cover_name": r.get("cover", ""), "start": start, "end": end,
            "old_url": r.get("old_url", ""), "social": social_links(r.get("links")),
            "offstage": any(d.get("offstage") for d in dirs_found),
            "rate_dir": next((d for d in dirs_found if d.get("offstage")), None) or next((d for d in dirs_found if d.get("criteria")), None),
        }
        if concert and concert["video"] and start is None:
            report.warn("Номера", row, f"У номера «{r['title']}» нет начала: видео откроется с самого начала записи.")
        numbers[nid.casefold()] = num
        concert["numbers"].append(num)

    # порядок внутри концерта
    for c in concerts.values():
        c["numbers"].sort(key=lambda n: (n["order"] is None, n["order"] or 0, n["start"] is None, n["start"] or 0, n["row"]))
        for i, n in enumerate(c["numbers"], start=1):
            n["position"] = i

    # Люди
    people = OrderedDict()
    extra = {norm(r["name"]): r for r in sheets["Люди"] if r.get("name")}

    def person(name):
        k = norm(name)
        if k not in people:
            e = extra.get(k, {})
            people[k] = {"name": name, "slug": p_slug[e["_row"]] if e else slugs_p.make(name), "numbers": [],
                         "about": e.get("about", ""), "link": e.get("link", "")}
        return people[k]

    for n in numbers.values():
        n["people"] = {"choreographers": [], "performers": []}
        for role_key, role_word in (("choreographers", "постановщик"), ("performers", "участник")):
            for name in n[role_key]:
                p = person(name)
                n["people"][role_key].append(p)
                entry = next((x for x in p["numbers"] if x["number"] is n), None)
                if entry:
                    entry["roles"].append(role_word)
                else:
                    p["numbers"].append({"number": n, "roles": [role_word]})
    for k, e in extra.items():
        if k not in people:
            report.warn("Люди", e["_row"], f"«{e['name']}» нет ни в одном номере. Строка пока ни на что не влияет.")
    loose = [p["name"] for k, p in people.items() if k not in extra]
    if loose:
        report.warn("Люди", None, f"У {len(loose)} человек нет строки на листе «Люди» (например, «{loose[0]}»): адрес их страницы зависит от написания имени. "
                                  "Выберите Каталог → Собрать список людей.")

    # похожие имена
    keys = list(people)
    by_tokens = defaultdict(list)
    for k in keys:
        by_tokens[" ".join(sorted(k.split()))].append(k)
    for group in by_tokens.values():
        if len(group) > 1:
            names = ", ".join(f"«{people[k]['name']}»" for k in group)
            report.warn("Номера", None, f"Возможно, это один человек: {names} (одинаковые слова в другом порядке). Приведите к одному написанию.")
    for i, a in enumerate(keys):
        for b in keys[i + 1:]:
            if abs(len(a) - len(b)) <= 2 and a[:2] == b[:2] and " ".join(sorted(a.split())) != " ".join(sorted(b.split())):
                if difflib.SequenceMatcher(None, a, b).ratio() >= 0.9:
                    report.warn("Номера", None, f"Возможно, это один человек: «{people[a]['name']}» и «{people[b]['name']}». Проверьте написание.")

    for p in people.values():
        def newest(x):
            c = x["number"]["concert"]
            when = c["date"].toordinal() if c["date"] else (dt.date(c["year"], 1, 1).toordinal() if c["year"] else 0)
            return (c["season_start"], when, -c["row"], x["number"]["position"])
        p["numbers"].sort(key=newest, reverse=True)
        for entry in p["numbers"]:
            entry["roles"] = sorted(set(entry["roles"]))
        p["is_director"] = any("постановщик" in e["roles"] for e in p["numbers"])
    return concerts, directions, numbers, people


# ---------------------------------------------------------------- обложки

COVER_MAX_BYTES = 8_388_608   # 8 МБ: больше обложку из ссылки не берём (формы принимают до 1,5 МБ, это проверяет скрипт)
_DL_CACHE = {}


def drive_direct(url):
    """Ссылка на файл Google Диска -> прямая ссылка на скачивание. Остальные ссылки без изменений."""
    m = re.search(r"drive\.google\.com/file/d/([\w-]+)", url) or re.search(r"drive\.google\.com/(?:open|uc)\?[^#]*?id=([\w-]+)", url)
    return f"https://drive.google.com/uc?export=download&id={m.group(1)}" if m else url


def download_cover(url, report, n):
    """Скачивает обложку по ссылке (Google Диск или любая прямая). Возвращает путь к файлу или None."""
    url = drive_direct(url.strip())
    if url in _DL_CACHE:
        return _DL_CACHE[url]
    path = None
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "socgum-catalog-build"})
        with urllib.request.urlopen(req, timeout=30) as r:
            data = r.read(COVER_MAX_BYTES + 1)
            ctype = r.headers.get("Content-Type", "")
        if len(data) > COVER_MAX_BYTES:
            report.warn("Номера", n["row"], f"Обложка по ссылке больше 8 МБ, не взята. Сожмите фото. Показана заглушка.")
        elif "image" not in ctype and not data[:4] in (b"\xff\xd8\xff\xe0", b"\xff\xd8\xff\xe1", b"\x89PNG", b"RIFF"):
            report.warn("Номера", n["row"], "По ссылке на обложку не картинка (проверьте, что файл открыт «всем, у кого есть ссылка»). Показана заглушка.")
        else:
            dl = Path(tempfile.gettempdir()) / "socgum-covers"
            dl.mkdir(parents=True, exist_ok=True)
            path = dl / (hashlib.sha1(url.encode()).hexdigest() + ".bin")
            path.write_bytes(data)
    except Exception as e:  # noqa: BLE001
        report.warn("Номера", n["row"], f"Не удалось скачать обложку по ссылке: {e}. Показана заглушка.")
    _DL_CACHE[url] = path
    return path


def find_cover(n, report=None):
    if re.match(r"https?://", n["cover_name"] or "", re.I):
        return download_cover(n["cover_name"], report, n) if report else None
    if not COVERS.exists():
        return None
    files = {f.name.casefold(): f for f in COVERS.iterdir() if f.is_file()}
    if n["cover_name"]:
        want = n["cover_name"].casefold()
        if want in files:
            return files[want]
        for ext in (".jpg", ".jpeg", ".png", ".webp"):
            if want + ext in files:
                return files[want + ext]
    for ext in (".jpg", ".jpeg", ".png", ".webp"):
        if n["id"].casefold() + ext in files:
            return files[n["id"].casefold() + ext]
    return None


def placeholder(seed):
    """Абстрактная «сценическая» картинка вместо отсутствующей обложки (без текста)."""
    rnd = random.Random(seed)
    w, h = 1200, 900
    base = Image.new("RGBA", (w, h), (13, 7, 2, 255))
    glow = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(glow)
    cx = rnd.randint(w // 4, 3 * w // 4)
    for _ in range(rnd.randint(4, 7)):
        spread = rnd.randint(60, 220)
        tx = rnd.randint(-200, w + 200)
        shade = rnd.choice([(255, 139, 54), (255, 176, 102), (214, 90, 24), (255, 213, 160)])
        d.polygon([(cx - 8, -20), (cx + 8, -20), (tx + spread, h + 20), (tx - spread, h + 20)],
                  fill=shade + (rnd.randint(28, 70),))
    glow = glow.filter(ImageFilter.GaussianBlur(26))
    floor = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    ImageDraw.Draw(floor).ellipse([w * 0.1, h * 0.82, w * 0.9, h * 1.15], fill=(255, 139, 54, 70))
    floor = floor.filter(ImageFilter.GaussianBlur(40))
    return Image.alpha_composite(Image.alpha_composite(base, glow), floor).convert("RGB")


def process_cover(img, outdir):
    outdir.mkdir(parents=True, exist_ok=True)
    img = ImageOps.exif_transpose(img).convert("RGB")
    card = ImageOps.fit(img, (640, 480), Image.LANCZOS, centering=(0.5, 0.35))
    card.save(outdir / "card.webp", "WEBP", quality=80, method=4)
    cover = img.copy()
    if cover.width > 1280:
        cover = cover.resize((1280, round(cover.height * 1280 / cover.width)), Image.LANCZOS)
    cover.save(outdir / "cover.webp", "WEBP", quality=82, method=4)
    og = ImageOps.fit(img, (1200, 630), Image.LANCZOS, centering=(0.5, 0.35))
    og.save(outdir / "og.jpg", "JPEG", quality=85, optimize=True, progressive=True)
    return cover.size


# ---------------------------------------------------------------- сборка сайта

def make_env():
    env = Environment(loader=FileSystemLoader(str(ROOT / "templates")),
                      autoescape=select_autoescape(["html"]), trim_blocks=True, lstrip_blocks=True)
    env.filters["hms"] = hms
    env.filters["vk_t"] = vk_t
    env.filters["ru_date"] = ru_date
    env.filters["domain"] = lambda u: re.sub(r"^https?://(www\.)?", "", str(u)).rstrip("/")
    return env


def write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def build_site(concerts, directions, numbers, people, report, demo, source):
    cfg = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))
    site_url = cfg.get("site_url", "").rstrip("/")
    base_path = urlparse(site_url).path.rstrip("/") + "/" if site_url else "/"
    if TMP.exists():
        shutil.rmtree(TMP)
    TMP.mkdir()
    shutil.copytree(ROOT / "assets", TMP / "assets")
    brand_font = None
    for ext in ("woff2", "woff"):
        if (ROOT / "assets" / "fonts" / f"brand.{ext}").exists():
            brand_font = ext
            write(TMP / "assets" / "brand.css",
                  '@font-face{font-family:"Brand";src:url(fonts/brand.%s) format("%s");font-weight:700;font-display:swap}\n'
                  ':root{--font-display:"Brand","Unbounded","Arial Black",sans-serif}\n' % (ext, ext))
            break

    # обложки
    for n in numbers.values():
        src = find_cover(n, report)
        if src:
            try:
                with Image.open(src) as im:
                    im.load()
                    size = process_cover(im, TMP / "img" / n["id"])
                n["has_cover"] = True
            except Exception as e:  # noqa: BLE001
                report.error("Номера", n["row"], f"Не удалось открыть обложку «{src.name}»: {e}")
                continue
        else:
            size = process_cover(placeholder(n["id"]), TMP / "img" / n["id"])
            n["has_cover"] = False
            if re.match(r"https?://", n["cover_name"] or "", re.I):
                pass  # причина уже в предупреждении о скачивании
            elif n["cover_name"]:
                report.warn("Номера", n["row"], f"Файл обложки «{n['cover_name']}» не найден в папке covers. Показана заглушка.")
            else:
                report.warn("Номера", n["row"], f"У номера «{n['title']}» нет обложки (положите в covers файл {n['id']}.jpg). Показана заглушка.")
        n["cover_size"] = size
    if report.errors:
        return None

    env = make_env()
    today = dt.date.today().isoformat()
    fbc = cfg.get("firebase") or {}
    firebase = {k: fbc[k] for k in ("apiKey", "authDomain", "projectId", "appId") if fbc.get(k)} if fbc.get("apiKey") and fbc.get("projectId") else None
    asset_v = hashlib.sha1(b"".join((ROOT / "assets" / f).read_bytes() for f in ("style.css", "app.js", "social.js"))).hexdigest()[:8]
    ctx_common = dict(firebase=firebase, asset_v=asset_v, site_title=cfg["site_title"], site_url=site_url, home_url=cfg.get("home_url", ""),
                      home_label=cfg.get("home_label", ""), demo=demo, brand_font=bool(brand_font),
                      lead=cfg.get("lead", ""), forms=list(FORMS))

    def concert_key(c):
        when = c["date"].toordinal() if c["date"] else (dt.date(c["year"], 1, 1).toordinal() if c["year"] else 0)
        return (-c["season_start"], -when, c["row"])

    concert_list = sorted(concerts.values(), key=concert_key)
    rank = {id(c): i for i, c in enumerate(concert_list)}
    all_numbers = sorted(numbers.values(), key=lambda n: (rank[id(n["concert"])], n["position"]))
    seasons = OrderedDict()
    for c in concert_list:
        if c["season"] and c["season"] not in seasons:
            seasons[c["season"]] = c["season_start"]
    used_dirs = [d for d in directions.values() if any(d in n["directions"] for n in numbers.values())]
    dirs_on = [d for d in used_dirs if not d["offstage"]]
    dirs_off = [d for d in used_dirs if d["offstage"]]
    all_people = sorted(people.values(), key=lambda p: norm(p["name"]))

    write(TMP / "index.html", env.get_template("index.html").render(
        root="", page_title=cfg["site_title"], numbers=all_numbers, concerts=concert_list, seasons=list(seasons),
        dirs_on=dirs_on, dirs_off=dirs_off, people=all_people, canonical=site_url + "/" if site_url else "",
        og_image="", description=cfg.get("lead", ""), **ctx_common))

    for n in numbers.values():
        people_names = ", ".join(p["name"] for p in n["people"]["choreographers"] + n["people"]["performers"])
        desc = n["about"] or (f"{n['collective']}. {n['concert']['name']}." if n["collective"] else n["concert"]["name"])
        write(TMP / "n" / n["id"] / "index.html", env.get_template("number.html").render(
            root="../../", n=n, page_title=f"{n['title']}: {cfg['site_title']}",
            canonical=f"{site_url}/n/{n['id']}/" if site_url else "",
            og_image=f"{site_url}/img/{n['id']}/og.jpg" if site_url else "", description=desc,
            video=None if n["offstage"] else n["concert"]["video"], watch_links=[] if n["offstage"] else n["concert"]["links"], people_names=people_names, **ctx_common))

    for p in people.values():
        write(TMP / "p" / p["slug"] / "index.html", env.get_template("person.html").render(
            root="../../", p=p, page_title=f"{p['name']}: {cfg['site_title']}",
            canonical=f"{site_url}/p/{p['slug']}/" if site_url else "", og_image="",
            description=f"Номера с участием: {p['name']}", **ctx_common))

    if firebase:
        write(TMP / "account" / "index.html", env.get_template("account.html").render(
            root="../", page_title=f"Кабинет: {cfg['site_title']}", canonical="", og_image="", description="", **ctx_common))
        write(TMP / "u" / "index.html", env.get_template("user.html").render(
            root="../", page_title=f"Пользователь: {cfg['site_title']}", canonical="", og_image="", description="", **ctx_common))

    write(TMP / "404.html", env.get_template("404.html").render(
        root=base_path, page_title=f"Страница не найдена: {cfg['site_title']}", canonical="", og_image="",
        description="", **ctx_common))

    feed = [{
        "id": n["id"], "title": n["title"], "collective": n["collective"],
        "concert": n["concert"]["name"], "concert_date": n["concert"]["date"].isoformat() if n["concert"]["date"] else None,
        "direction": n["direction"]["name"] if n["direction"] else None,
        "directions": [d["name"] for d in n["directions"]], "position": n["position"],
        "start": n["start"], "end": n["end"], "performers": n["performers"], "choreographers": n["choreographers"],
        "old_url": n["old_url"], "url": f"n/{n['id']}/"} for n in all_numbers]
    people_feed = [{"id": p["slug"], "name": p["name"], "director": p["is_director"]}
                   for p in sorted(people.values(), key=lambda p: p["slug"])]
    for item, n in zip(feed, all_numbers):
        item["director_ids"] = [p["slug"] for p in n["people"]["choreographers"]]
    # отпечаток того, что нужно сервису комментариев: меняется, только когда меняются люди или постановщики
    sync_hash = hashlib.sha1(json.dumps(
        [people_feed, sorted((f["id"], f["title"], f["director_ids"]) for f in feed)], ensure_ascii=False).encode("utf-8")).hexdigest()[:16]
    write(TMP / "catalog.json", json.dumps({"generated": today, "sync_hash": sync_hash, "people": people_feed, "numbers": feed},
                                           ensure_ascii=False, indent=1))

    if site_url:
        urls = [f"{site_url}/"] + [f"{site_url}/n/{n['id']}/" for n in numbers.values()] + [f"{site_url}/p/{p['slug']}/" for p in people.values()]
        write(TMP / "sitemap.xml", '<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
              + "".join(f"<url><loc>{html.escape(u)}</loc><lastmod>{today}</lastmod></url>\n" for u in urls) + "</urlset>\n")
        write(TMP / "robots.txt", f"User-agent: *\nAllow: /\nSitemap: {site_url}/sitemap.xml\n")
    write(TMP / ".nojekyll", "")
    return TMP


def main():
    report = Report()
    wb, source, demo = open_workbook(report)
    stats = ""
    ok = False
    if wb is not None:
        concerts, directions, numbers, people = build_model(wb, report)
        stats = (f"Итог: концертов {len(concerts)}, номеров {len(numbers)}, людей {len(people)}.")
        if not report.errors:
            built = build_site(concerts, directions, numbers, people, report, demo, source)
            if built and not report.errors:
                if OUT.exists():
                    shutil.rmtree(OUT)
                built.rename(OUT)
                ok = True
    if not ok and TMP.exists():
        shutil.rmtree(TMP)
    md = report.markdown(source, stats)
    (ROOT / "build-report.md").write_text(md, encoding="utf-8")
    print(md)
    summary = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary:
        with open(summary, "a", encoding="utf-8") as f:
            f.write(md + "\n")
    if os.environ.get("GITHUB_ACTIONS"):
        for where, text in report.errors:
            print(f"::error title=Ошибка в таблице ({where})::{text}")
        for where, text in report.warnings[:30]:
            print(f"::warning title=Проверьте ({where})::{text}")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
